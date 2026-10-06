"""Gera as planilhas FICTÍCIAS usadas nos testes (nunca use dados reais da empresa nos testes).
Uso: python3 tests/gerar_fixtures.py   → cria os arquivos em fixtures/
"""
import csv, io, random, datetime as dt
from pathlib import Path
import openpyxl

OUT = Path(__file__).resolve().parent.parent / "fixtures"
OUT.mkdir(exist_ok=True)
D0 = dt.datetime(2026, 4, 1)

def fluxo(n, k, inicio, passo=(1, 3)):
    """datas acumuladas até a coluna k (as colunas seguintes ficam vazias)"""
    out, d = [], inicio
    for i in range(n):
        if i > k: out.append(None); continue
        if i: d = d + dt.timedelta(days=random.randint(*passo))
        out.append(d)
    return out

def times():
    """Portfólio com 7 times de fluxos diferentes (ilhas, fileiras, colunas longas)."""
    random.seed(4)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    I = wb.create_sheet("Iniciativa"); I.append(["ID","Title","AnoSemestreRoadmap","Assigned To","Materialização da Oportunidade ou Solicitação","Em Execução","Concluído"])
    I.append([1, "[Cartões] Automatizar processo de Disputa", "2026 2º Semestre", "Fulano de Tal <fulano@exemplo.com>", D0, D0, None])
    R = wb.create_sheet("Release"); R.append(["ID","Title","Parent","Assigned To","Inventário de Opções de Valor","Em Desenvolvimento","Entregue"])
    E = wb.create_sheet("Épico"); E.append(["ID","Title","Parent","Target Date","Backlog","Em Desenvolvimento","Fechado"])
    for r in range(3):
        R.append([10 + r, f"Release {r}", 1, "Ciclano", D0, D0, None])
        for e in range(4):
            E.append([100 + r * 10 + e, f"Épico {r}-{e}", 10 + r, dt.datetime(2026, 8, 1), D0, D0, D0 if e % 2 else None])
    base = ["Backlog","Refinamento","READY / PRONTO PARA DEV","Em Desenvolvimento","Code Review","Testes","Pronto para Deploy","Em Produção","Fechado"]
    fl = {"TIME CORE": base, "TIME IB": base[:3] + ["UX-UI Doing","UX-UI Done"] + base[3:], "TIME BO": base, "TIME MOBILE": base[:4] + ["QA"] + base[4:],
          "TIME PAGAMENTOS": base, "TIME DADOS": base[:2] + ["Modelagem"] + base[2:], "TIME CANAIS": base}
    oid = 1000
    for nome, cols in fl.items():
        T = wb.create_sheet(nome); T.append(["ID","ID_EPICO_UNICRED","Title","Work Item Type","Tags"] + cols)
        for _ in range(random.randint(8, 30)):
            k = len(cols) - 1 if random.random() < .7 else random.randint(0, len(cols) - 1)
            oid += 1
            T.append([oid, 100 + random.choice([0,1,2,3,10,11,12,13,20,21,22,23]), f"Item {oid}", random.choice(["User Story","Technical Story","Internal Bug"]),
                      "[Blocked]" if oid % 17 == 0 else ""] + fluxo(len(cols), k, D0 + dt.timedelta(days=random.randint(0, 60))))
    wb.save(OUT / "times.xlsx")

def responsaveis():
    """300 iniciativas e ~280 responsáveis (filtro com busca e múltipla escolha)."""
    random.seed(9)
    nomes = ["Ana","Bruno","Carla","Diego","Elisa","Fábio","Gustavo","Helena","Igor","Júlia","Karen","Lucas","Marina","Nícolas","Otávio","Paula"]
    sobren = ["Silva","Souza","Guzzo","Conceição","Pereira","Araújo","Lima","Schmidt"]
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    I = wb.create_sheet("Iniciativa"); I.append(["ID","Title","AnoSemestreRoadmap","Assigned To","Materialização da Oportunidade ou Solicitação","Concluído"])
    R = wb.create_sheet("Release"); R.append(["ID","Title","Parent","Assigned To","Inventário de Opções de Valor","Entregue"])
    E = wb.create_sheet("Épico"); E.append(["ID","Title","Parent","Target Date","Backlog","Fechado"])
    T = wb.create_sheet("TIME CORE"); T.append(["ID","ID_EPICO_UNICRED","Title","Work Item Type","Backlog","READY / PRONTO PARA DEV","Pronto para Deploy","Fechado"])
    for i in range(300):
        quem = None if i % 37 == 0 else f"{random.choice(nomes)} {random.choice(nomes)} {random.choice(sobren)} - Empresa <p{i}@exemplo.com>"
        I.append([100 + i, f"Iniciativa {i}", "2026 2º Semestre", quem, D0, None])
        R.append([5000 + i, f"Rel {i}", 100 + i, "Outra Pessoa", D0, None])
        E.append([7000 + i, f"Ep {i}", 5000 + i, D0, D0, None])
    wb.save(OUT / "responsaveis.xlsx")

def fluxo_largo():
    """Um time com 38 colunas e itens concentrados no fim (etapas vazias ocultas, enquadramento)."""
    random.seed(1)
    cols = ["Backlog","Próximos","Resultado de Negócio (métrica)","Detalhamento","Em Detalhamento","Detalhamento Doing","Detalhamento Done","Disponível para UX",
            "UX-UI Doing","UX-UI Done","Pré Refinamento","Aguard. Arquitetura de Solução","Arquitetura de Solução","Pronto para Refinamento","Refinamento Doing",
            "Refinamento Done","Design Técnico","Aguard. Mesa Arquitetura","Pronto para Refinar","Refinamento","Pronto para Design","Refinando Doing","Refinando Done",
            "Refinado","Aguard. Design Técnico","Design Técnico Doing","Design Técnico Done","READY / PRONTO PARA DEV","Em Desenv.","Aguard. Review","Aguard. Ambiente",
            "Desenvolvimento Doing","Desenvolvimento Done","Disponível para Teste","Testes","Pronto para Deploy","Em Produção","Fechado"]
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    I = wb.create_sheet("Iniciativa"); I.append(["ID","Title","AnoSemestreRoadmap","Materialização da Oportunidade ou Solicitação","Análise","Concluído"]); I.append([611510,"Nova Fatura","2026 2º Semestre",D0,None,None])
    R = wb.create_sheet("Release"); R.append(["ID","Title","Parent","Assigned To","Inventário de Opções de Valor","Planejada","Entregue"]); R.append([747178,"Fatura V2",611510,"X",D0,D0,None])
    E = wb.create_sheet("Épico"); E.append(["ID","Title","Parent","Target Date","Backlog","Fechado"]); E.append([498565,"Gamificação da fatura",747178,dt.datetime(2026,6,15),D0,D0])
    T = wb.create_sheet("TIME IB"); T.append(["ID","ID_EPICO_UNICRED","Title","Work Item Type"] + cols)
    for i in range(11):
        T.append([900 + i, 498565, f"Item operacional {i + 1}", "User Story"] + fluxo(len(cols), random.choice([35, 36, 37]), D0, (1, 1)))
    wb.save(OUT / "fluxo_largo.xlsx")

def desdobramento():
    """Iniciativas/releases com e sem desdobramento, abertas e concluídas (regra B)."""
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    I = wb.create_sheet("Iniciativa"); I.append(["ID","Title","AnoSemestreRoadmap","Materialização da Oportunidade ou Solicitação","Concluído"])
    I.append([1, "Ini completa e fechada", "2026 1º Semestre", D0, None])
    I.append([2, "Ini com algo aberto", "2026 1º Semestre", D0, None])
    I.append([3, "Ini sem release (aberta)", "2026 1º Semestre", D0, None])
    I.append([4, "Ini sem release (concluída)", "2026 1º Semestre", D0, D0])
    R = wb.create_sheet("Release"); R.append(["ID","Title","Parent","Inventário de Opções de Valor","Entregue"])
    R.append([10, "Rel fechada", 1, D0, None]); R.append([20, "Rel fechada B", 2, D0, None]); R.append([21, "Rel aberta", 2, D0, None])
    R.append([22, "Rel sem épico (aberta)", 2, D0, None]); R.append([23, "Rel sem épico (entregue)", 2, D0, D0]); R.append([99, "Rel órfã", 999, D0, None])
    E = wb.create_sheet("Épico"); E.append(["ID","Title","Parent","Target Date","Backlog","Fechado"])
    for eid, par in [(100, 10), (101, 10), (200, 20), (210, 21)]: E.append([eid, f"Ep {eid}", par, D0, D0, None])
    T = wb.create_sheet("TIME CORE"); T.append(["ID","ID_EPICO_UNICRED","Title","Work Item Type","Backlog","READY / PRONTO PARA DEV","Em Desenvolvimento","Pronto para Deploy","Fechado"])
    oid = 1
    for eid in (100, 101, 200):
        for _ in range(3): T.append([oid, eid, "x", "User Story", D0, D0, D0, D0, D0]); oid += 1
    T.append([oid, 210, "x", "User Story", D0, D0, D0, None, None]); oid += 1
    T.append([oid, 210, "x", "User Story", D0, D0, D0, D0, D0])
    wb.save(OUT / "desdobramento.xlsx")


def f4p():
    """Time único (CORE) com 10 itens de CT conhecido, para conferir o P95/P50 do Report F4P
    (docs/backlog/report-f4p.md) contra um cálculo independente. Datas relativas a agora,
    para a amostra sempre cair dentro dos últimos 6 meses (padrão de CFG.f4p.months)."""
    entry = dt.datetime.now() - dt.timedelta(days=150)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    I = wb.create_sheet("Iniciativa"); I.append(["ID","Title","AnoSemestreRoadmap","Materialização da Oportunidade ou Solicitação","Concluído"])
    I.append([1, "Ini F4P", "2026 2º Semestre", entry, None])
    R = wb.create_sheet("Release"); R.append(["ID","Title","Parent","Inventário de Opções de Valor","Entregue"]); R.append([10, "Rel F4P", 1, entry, None])
    E = wb.create_sheet("Épico"); E.append(["ID","Title","Parent","Target Date","Backlog","Fechado"]); E.append([100, "Ep F4P", 10, entry, entry, None])
    T = wb.create_sheet("TIME CORE"); T.append(["ID","ID_EPICO_UNICRED","Title","Work Item Type","Backlog","READY / PRONTO PARA DEV","Pronto para Deploy","Fechado"])
    for i, ct in enumerate([10, 20, 25, 30, 35, 40, 45, 50, 60, 100]):        # CTs conhecidos (n=10)
        deploy = entry + dt.timedelta(days=ct)
        T.append([2000 + i, 100, f"Item {i}", "User Story", entry, entry, deploy, deploy + dt.timedelta(days=1)])
    wb.save(OUT / "f4p.xlsx")

def relatorio():
    """Quatro times (CORE, IB, BO, MOBILE) com itens reservados (tag ROADMAP) no roadmap vigente — alimenta
    o gerador do relatório por e-mail (scripts/relatorio) e seu teste. Datas relativas a agora, para o
    CycleTime e o dead line caírem sempre numa situação realista, qualquer que seja o dia da execução."""
    random.seed(11)
    hoje = dt.datetime.now(); ini = hoje - dt.timedelta(days=150)
    sem = f"{hoje.year} {1 if hoje.month <= 6 else 2}º Semestre"
    alvo = dt.datetime(hoje.year, 3 if hoje.month <= 6 else 9, 15)    # Target Date dentro do semestre vigente
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    I = wb.create_sheet("Iniciativa"); I.append(["ID","Title","AnoSemestreRoadmap","Assigned To","Materialização da Oportunidade ou Solicitação","Concluído"])
    I.append([1, "[Cartões] Iniciativa fictícia A", sem, "Fulano de Tal <fulano@exemplo.com>", ini, None])
    I.append([2, "[Cartões] Iniciativa fictícia B", sem, "Ciclana de Tal <ciclana@exemplo.com>", ini, None])
    R = wb.create_sheet("Release"); R.append(["ID","Title","Parent","Assigned To","Inventário de Opções de Valor","Entregue"])
    E = wb.create_sheet("Épico"); E.append(["ID","Title","Parent","Target Date","Backlog","Fechado"])
    epicos = []
    for r in range(4):
        R.append([10 + r, f"Release fictícia {r}", 1 + r % 2, "Ciclano", ini, None])
        for e in range(3):
            eid = 100 + r * 10 + e; epicos.append(eid)
            E.append([eid, f"Épico fictício {r}-{e}", 10 + r, alvo + dt.timedelta(days=30 * e), ini, None])
    base = ["Backlog","Refinamento","READY / PRONTO PARA DEV","Em Desenvolvimento","Code Review","Testes","Pronto para Deploy","Em Produção","Fechado"]
    oid = 3000
    for nome in ("TIME CORE", "TIME IB", "TIME BO", "TIME MOBILE"):
        T = wb.create_sheet(nome); T.append(["ID","ID_EPICO_UNICRED","Title","Work Item Type","Tags"] + base)
        for _ in range(random.randint(18, 26)):
            k = len(base) - 1 if random.random() < .35 else random.randint(0, len(base) - 2)
            oid += 1
            tags = [t for t, p in (("ROADMAP", .6), ("Blocked", .08)) if random.random() < p]
            T.append([oid, random.choice(epicos), f"Item fictício {oid}", random.choice(["User Story","User Story","Technical Story","Internal Bug"]),
                      "; ".join(tags)] + fluxo(len(base), k, ini + dt.timedelta(days=random.randint(0, 70)), passo=(2, 9)))
    wb.save(OUT / "relatorio.xlsx")

if __name__ == "__main__":
    for f in (times, responsaveis, fluxo_largo, desdobramento, f4p, relatorio): f()
    print("fixtures geradas em", OUT)
