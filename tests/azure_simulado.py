"""Azure DevOps simulado para os testes, gerado a partir de fixtures/times.xlsx (dados fictícios).
Reproduz os cenários reais encontrados na integração:
- duas organizações: 'org-portfolio' (Iniciativa, Release, Épico e o time DADOS) e 'org-times' (demais times);
- times ligados ao épico pelo campo ID_EPICO_UNICRED (org-times) e pelo Parent (DADOS, na org-portfolio);
- quadros com fluxos próprios (o de épicos termina em 'Fechado'; o do DADOS em 'Pronto');
- um item no estado Removed (deve ser excluído), uma coluna antiga no histórico (ignorada por padrão)
  e o token 'token-errado' (401)."""
import json, re, base64, datetime as dt
from urllib.parse import urlsplit, parse_qs, unquote
import openpyxl
from conftest import FIX

CORS = {"Access-Control-Allow-Origin": "*", "Access-Control-Allow-Headers": "authorization,content-type", "Access-Control-Allow-Methods": "GET,POST,OPTIONS"}
TOKEN_RUIM = "Basic " + base64.b64encode(b":token-errado").decode()

def _dia(v): return v.strftime("%Y-%m-%d") if isinstance(v, dt.datetime) else (str(v)[:10] if v else "")

def _colunas(nomes, prefixo):
    cols, keys, i = [], [], 0
    while i < len(nomes):
        n = nomes[i]
        if n.endswith(" Doing") and i + 1 < len(nomes) and nomes[i + 1] == n[:-6] + " Done":
            cid = f"{prefixo}-{i}"; cols.append({"id": cid, "name": n[:-6], "isSplit": True}); keys += [(n, cid, "Doing"), (nomes[i + 1], cid, "Done")]; i += 2
        else:
            cid = f"{prefixo}-{i}"; cols.append({"id": cid, "name": n, "isSplit": False}); keys.append((n, cid, "Unknown")); i += 1
    return cols, keys

def _revisoes(iid, datas, keys):
    """uma revisão por mudança de coluna: grupo de datas iguais = entrou na última coluna do grupo"""
    rv, n = [], 0
    criado = next((d for d in datas if d), "2026-01-01")
    for j, d in enumerate(datas):
        if d and (j + 1 == len(datas) or datas[j + 1] != d):
            n += 1
            nome, cid, done = keys[j]
            rv.append({"WorkItemId": iid, "Revision": n, "ChangedDate": f"{d}T12:{n:02d}:00-03:00", "CreatedDate": f"{criado}T09:00:00-03:00",
                       "State": "Active", "TagNames": None,
                       "BoardLocations": [{"ColumnId": cid, "ColumnName": nome.replace(" Doing", "").replace(" Done", ""), "Done": done, "LaneName": "Default Lane"}]})
    return rv

class AzureSimulado:
    def __init__(self):
        wb = openpyxl.load_workbook(FIX / "times.xlsx")
        self.quadros = {}   # (org, projeto, time, nível) -> dados do quadro
        self.itens, self.revs, self.rels = {}, {}, {}
        def ler(ws):
            cab = [c.value for c in ws[1]]; return cab, [dict(zip(cab, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
        def quadro(org, proj, time, nivel, cab, linhas, ini, fim, tipos, campos):
            fl = cab[cab.index(ini):cab.index(fim) + 1]
            cols, keys = _colunas(fl, f"{time}-{nivel}")
            ids = []
            for r in linhas:
                iid = int(r["ID"]); ids.append(iid)
                self.revs[(org, iid)] = _revisoes(iid, [_dia(r.get(k)) for k in fl], keys)
                f = {"System.Id": iid, "System.Title": r["Title"], "System.WorkItemType": r.get("Work Item Type") or tipos[0], "System.State": "Active"}
                f.update(campos(r)); self.itens[(org, iid)] = f
            self.quadros[(org, proj, time, nivel)] = {"cols": cols, "ids": ids, "tipos": tipos, "area": f"{proj}\\{time}"}
        c, l = ler(wb["Iniciativa"]); quadro("org-portfolio", "Portfolio", "Portfolio UBR", "Iniciativas", c, l, "Materialização da Oportunidade ou Solicitação", "Concluído", ["Initiative"],
            lambda r: {"Custom.AnoSemestreRoadmap": r["AnoSemestreRoadmap"], "System.AssignedTo": {"displayName": "Fulano de Tal", "uniqueName": "fulano@exemplo.com"}})
        c, l = ler(wb["Release"]); quadro("org-portfolio", "Portfolio", "Portfolio UBR", "Releases", c, l, "Inventário de Opções de Valor", "Entregue", ["Product Release"],
            lambda r: {"System.Parent": r["Parent"]})
        c, l = ler(wb["Épico"]); quadro("org-portfolio", "Portfolio", "Coordenacao Epicos", "Epicos", c, l, "Backlog", "Fechado", ["Epic"],
            lambda r: {"System.Parent": r["Parent"], "Microsoft.VSTS.Scheduling.TargetDate": _dia(r["Target Date"]) + "T03:00:00Z"})
        for nome in [n for n in wb.sheetnames if n.startswith("TIME ")]:
            c, l = ler(wb[nome]); time = nome[5:]
            if time == "DADOS":   # time na organização do portfólio: vínculo pelo Parent; fluxo termina em "Pronto"
                c = [("Pronto" if x == "Fechado" else x) for x in c]
                l = [{("Pronto" if k == "Fechado" else k): v for k, v in r.items()} for r in l]
                quadro("org-portfolio", "TI", "DADOS", "Stories", c, l, "Backlog", "Pronto", ["User Story", "Technical Story", "Internal Bug"],
                       lambda r: {"System.Parent": r["ID_EPICO_UNICRED"], "System.Tags": (r.get("Tags") or "").strip("[]")})
            else:
                quadro("org-times", "Times", time, "Stories", c, l, "Backlog", "Fechado", ["User Story", "Technical Story", "Internal Bug"],
                       lambda r: {"Custom.ID_EPICO_UNICRED": r["ID_EPICO_UNICRED"], "System.Tags": (r.get("Tags") or "").strip("[]")})
        # item removido (só existe no Azure) e coluna antiga no histórico de um item do CORE
        core = self.quadros[("org-times", "Times", "CORE", "Stories")]
        self.itens[("org-times", 99999)] = {"System.Id": 99999, "System.Title": "Removido", "System.WorkItemType": "User Story", "System.State": "Removed"}
        self.revs[("org-times", 99999)] = []; core["ids"].append(99999)
        primeiro = core["ids"][0]
        self.revs[("org-times", primeiro)].insert(0, {"WorkItemId": primeiro, "Revision": 0, "ChangedDate": "2025-01-01T10:00:00-03:00", "CreatedDate": "2025-01-01T09:00:00-03:00",
            "State": "New", "TagNames": None, "BoardLocations": [{"ColumnId": "coluna-antiga", "ColumnName": "Coluna Antiga", "Done": "Unknown", "LaneName": "x"}]})
        self.chamadas = 0

    def rota(self, route):
        req = route.request; u = unquote(req.url); sp = urlsplit(u); path = sp.path
        if req.method == "OPTIONS": return route.fulfill(status=200, headers=CORS, body="")
        js = lambda o, st=200: route.fulfill(status=st, headers={**CORS, "Content-Type": "application/json"}, body=json.dumps(o))
        if req.headers.get("authorization") == TOKEN_RUIM: return js({"message": "unauthorized"}, 401)
        self.chamadas += 1
        seg = path.strip("/").split("/"); org = seg[0]
        if sp.netloc.startswith("analytics"):
            if path.endswith("/Projects"): return js({"value": [{"ProjectName": "x"}]})
            ids = [int(x) for x in re.findall(r"\d+", parse_qs(sp.query)["$filter"][0])]
            return js({"value": [v for i in ids for v in self.revs.get((org, i), [])]})
        projs = sorted({k[1] for k in self.quadros if k[0] == org})
        if path.endswith("/_apis/projects"): return js({"value": [{"id": p, "name": p} for p in projs]})
        if path.endswith("/teams"): return js({"value": [{"name": t} for t in sorted({k[2] for k in self.quadros if k[0] == org and k[1] == seg[3]})]})
        if path.endswith("/wit/fields"):
            return js({"value": [{"name": "AnoSemestreRoadmap", "referenceName": "Custom.AnoSemestreRoadmap"}] +
                       ([{"name": "ID EPICO UNICRED", "referenceName": "Custom.ID_EPICO_UNICRED"}] if org == "org-times" else [])})
        if "/_apis/work/" in path:
            proj, time = seg[1], seg[2]
            qs = {k[3]: v for k, v in self.quadros.items() if k[:3] == (org, proj, time)}
            if path.endswith("/work/backlogs"): return js({"value": [{"name": n, "rank": 5 - i, "workItemTypes": [{"name": t} for t in q["tipos"]]} for i, (n, q) in enumerate(qs.items())]})
            if path.endswith("teamfieldvalues"): return js({"values": [{"value": next(iter(qs.values()))["area"], "includeChildren": False}]})
            if "/work/boards/" in path: return js({"value": qs[path.split("/work/boards/")[1].split("/")[0]]["cols"]})
        if path.endswith("/wit/wiql"):
            q = json.loads(req.post_data)["query"]
            alvo = next(v for k, v in self.quadros.items() if k[0] == org and f"'{v['area']}'" in q and f"'{v['tipos'][0]}'" in q)
            return js({"workItems": [{"id": i} for i in alvo["ids"]]})
        if path.endswith("/wit/workitemsbatch"):
            b = json.loads(req.post_data)
            if b.get("$expand"): return js({"value": [{"id": i, "relations": []} for i in b["ids"]]})
            return js({"value": [{"id": i, "fields": {k: v for k, v in self.itens[(org, i)].items() if k in b["fields"]}} for i in b["ids"] if (org, i) in self.itens]})
        return route.fulfill(status=404, headers=CORS, body="não simulado: " + path)

FONTES = [
    {"id": "s1", "role": "ini", "org": "org-portfolio", "project": "Portfolio", "team": "Portfolio UBR", "level": "Iniciativas", "alias": ""},
    {"id": "s2", "role": "rel", "org": "org-portfolio", "project": "Portfolio", "team": "Portfolio UBR", "level": "Releases", "alias": ""},
    {"id": "s3", "role": "epi", "org": "org-portfolio", "project": "Portfolio", "team": "Coordenacao Epicos", "level": "Epicos", "alias": ""},
    {"id": "s4", "role": "op", "org": "org-portfolio", "project": "TI", "team": "DADOS", "level": "Stories", "alias": "DADOS"}] + [
    {"id": f"t{n}", "role": "op", "org": "org-times", "project": "Times", "team": n, "level": "Stories", "alias": n}
    for n in ["CORE", "IB", "BO", "MOBILE", "PAGAMENTOS", "CANAIS"]]
